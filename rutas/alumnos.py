"""
Búsqueda de alumnos, alta individual/masiva, expediente, ficha de
inscripción, cambio de estatus y avance de cuatrimestre.

ver_expediente y ficha_inscripcion llevan @rol_requerido('DIRECTIVO',
'ADMINISTRATIVO', 'CAPTURADOR') -- fix #1 de la auditoría de producción
(antes solo tenían @login_required, sin restricción de rol). No
debilitar este decorador al mover estas dos rutas.
"""

import io
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from flask import Blueprint, render_template, request, flash, redirect, url_for, send_file
from flask_login import login_required, current_user
from sqlalchemy import or_
from sqlalchemy.exc import IntegrityError, SQLAlchemyError

from extensiones import db
from modelos import (
    Alumno, EstatusAlumno, PlanEstudio, HistorialEstatus,
    Materia, DocumentoAlumno, TipoDocumento, TurnoAlumno, ModalidadEstudio,
)
from utilidades.seguridad import rol_requerido
from utilidades.archivos import (
    CARACTERES_FORMULA_EXCEL, error_de_tamano_xlsx, error_de_dimensiones_hoja,
)
from utilidades.fechas import ahora_utc
from utilidades.paginacion import ALUMNOS_POR_PAGINA, pagina_valida
from servicios.alumnos import calcular_estadisticas_alumnos, _matriculas_con_adeudo
from servicios.auditoria import registrar
from servicios.academico import _avanzar_cuatrimestre, _max_periodos, _generar_carga_academica
from rutas.registro import CORREO_REGEX
from servicios.matriculas import crear_alumno_generando_matricula
from servicios.cobros import _generar_cargos_de_inscripcion, _generar_cargos_de_reinscripcion

alumnos_bp = Blueprint('alumnos', __name__)


@alumnos_bp.route('/')
@login_required
def index():
    """
    Pantalla de inicio: dashboard con números clave + filtros rápidos,
    y el buscador universal. Si viene ?filtro=algo en la URL, muestra
    esa lista filtrada en vez del dashboard vacío.
    """
    # Quién debe dinero es información de Cobros: un CAPTURADOR no la ve (ni la tarjeta ni el filtro).
    ve_cobros = current_user.puede_cobros()
    estadisticas = calcular_estadisticas_alumnos(incluir_adeudo=ve_cobros)

    filtros_disponibles = {
        'pendientes': ('Alumnos Pendientes de Validación', Alumno.estatus == EstatusAlumno.PENDIENTE),
        'activos': ('Alumnos Activos', Alumno.estatus == EstatusAlumno.ACTIVO),
        'documentacion': ('Alumnos con Documentación Pendiente', Alumno.documentacion_pendiente.isnot(None)),
        'faltas': ('Alumnos con Faltas Administrativas', Alumno.faltas_administrativas.isnot(None)),
        'con_adeudo': ('Alumnos con Adeudo Económico', None),
        'recientes': ('Últimos 10 Alumnos Registrados', None),
    }

    filtro = request.args.get('filtro')
    if filtro == 'con_adeudo' and not ve_cobros:
        filtro = None
    termino = request.args.get('q', '').strip()
    page = pagina_valida(request.args.get('page', 1, type=int))
    resultados = None
    titulo_filtro = None
    paginacion = None

    if termino:
        # Búsqueda por texto. Llega aquí como GET (ver buscar() abajo) para
        # que los enlaces de paginación puedan conservar el término.
        consulta = Alumno.query.filter(
            or_(
                Alumno.matricula_id == termino,
                Alumno.nombre_completo.ilike(f'%{termino}%')
            )
        ).order_by(Alumno.nombre_completo.asc())

        paginacion = consulta.paginate(page=page, per_page=ALUMNOS_POR_PAGINA, error_out=False)
        resultados = paginacion.items

        if not resultados and page == 1:
            flash(f'No se encontraron alumnos que coincidan con "{termino}".', 'danger')

    elif filtro in filtros_disponibles:
        titulo_filtro, condicion = filtros_disponibles[filtro]

        if filtro == 'recientes':
            # Ya viene acotado a 10 por definición: no necesita paginarse.
            resultados = Alumno.query.order_by(Alumno.fecha_registro.desc()).limit(10).all()

        elif filtro == 'con_adeudo':
            matriculas = _matriculas_con_adeudo()
            if matriculas:
                consulta = (
                    Alumno.query
                    .filter(Alumno.matricula_id.in_(matriculas))
                    .order_by(Alumno.nombre_completo.asc())
                )
                paginacion = consulta.paginate(page=page, per_page=ALUMNOS_POR_PAGINA, error_out=False)
                resultados = paginacion.items
            else:
                resultados = []

        else:
            consulta = Alumno.query.filter(condicion).order_by(Alumno.nombre_completo.asc())
            paginacion = consulta.paginate(page=page, per_page=ALUMNOS_POR_PAGINA, error_out=False)
            resultados = paginacion.items

    return render_template(
        'buscador.html',
        estadisticas=estadisticas,
        resultados=resultados,
        titulo_filtro=titulo_filtro,
        termino=termino or None,
        filtro=filtro,
        paginacion=paginacion
    )


@alumnos_bp.route('/buscar', methods=['POST'])
@login_required
def buscar():
    """
    Búsqueda flexible tipo 'banco':
      - Coincidencia EXACTA por matricula_id.
      - O coincidencia PARCIAL (case-insensitive) por nombre_completo.
    """
    termino = request.form.get('termino', '').strip()

    if not termino:
        flash('Ingresa una matrícula o un nombre para buscar.', 'warning')
        return redirect(url_for('alumnos.index'))

    # La búsqueda en sí vive en index(): aquí solo se redirige pasando el
    # término en la URL. Dos motivos:
    #   1. Los enlaces "Siguiente/Anterior" de la paginación son GET; si el
    #      resultado se renderizara aquí (POST), al pasar de página se
    #      perdería el término buscado.
    #   2. La URL queda compartible y se puede recargar sin que el
    #      navegador pregunte "¿reenviar formulario?".
    return redirect(url_for('alumnos.index', q=termino))


# (columna, es_obligatoria, descripción para la hoja de ayuda de la plantilla)
COLUMNAS_IMPORTACION_ALUMNOS = [
    ('nombre_completo', True, 'Nombre completo del alumno'),
    ('curp', True, 'CURP, 18 caracteres'),
    ('fecha_nacimiento', True, 'Formato AAAA-MM-DD, ej. 2005-03-21'),
    ('fecha_certificado_prepa', True, 'Formato AAAA-MM-DD'),
    ('clave_carrera', True, 'Clave del plan de estudios (ver hoja "Guía" para las claves válidas)'),
    ('sexo', True, 'Femenino / Masculino / Otro'),
    ('estatus', False, 'PENDIENTE / ACTIVO / BAJA_TEMPORAL / BAJA_DEFINITIVA / EGRESADO (vacío = ACTIVO)'),
    ('cuatrimestre_actual', False, 'Número de cuatrimestre en el que va (vacío = 1)'),
    ('correo', False, ''),
    ('telefono', False, 'Teléfono fijo'),
    ('telefono_movil', False, ''),
    ('domicilio_calle_numero', False, ''),
    ('domicilio_ciudad', False, ''),
    ('domicilio_cp', False, ''),
    ('domicilio_estado', False, ''),
    ('numero_identificacion', False, 'INE u otra identificación oficial'),
    ('estado_civil', False, ''),
    ('nacionalidad', False, 'Vacío = "Mexicana"'),
    ('tipo_sangre', False, 'Ej. O+, A-'),
    ('contacto_emergencia_nombre', False, ''),
    ('contacto_emergencia_telefono', False, ''),
    ('contacto_emergencia_parentesco', False, 'Ej. Madre, Hermana'),
    ('turno', False, 'MATUTINO / VESPERTINO / MIXTO'),
    ('modalidad', False, 'ESCOLARIZADO / SEMIESCOLARIZADO / DISTANCIA'),
    ('grupo_actual', False, ''),
]


# Longitud máxima de cada columna de texto, leída del modelo: SQLite no la hace
# cumplir, PostgreSQL sí (y revienta con 500 a mitad del lote).
LONGITUDES_IMPORTACION = {
    nombre: Alumno.__table__.c[nombre].type.length
    for nombre, _, _ in COLUMNAS_IMPORTACION_ALUMNOS
    if nombre in Alumno.__table__.c and getattr(Alumno.__table__.c[nombre].type, 'length', None)
}


def _texto_de_celda(valor):
    """Excel entrega los teléfonos como número (5512345678.0): se guardan como texto sin el ".0"."""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


@alumnos_bp.route('/alumnos/importar/plantilla')
@rol_requerido('DIRECTIVO')
def plantilla_importacion():
    """Genera y descarga el archivo .xlsx en blanco con las columnas esperadas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Alumnos'

    fuente_encabezado = Font(bold=True, color='FFFFFF')
    relleno_obligatorio = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    relleno_opcional = PatternFill(start_color='6C757D', end_color='6C757D', fill_type='solid')

    for col_idx, (nombre_col, requerido, _desc) in enumerate(COLUMNAS_IMPORTACION_ALUMNOS, start=1):
        celda = ws.cell(row=1, column=col_idx, value=nombre_col)
        celda.font = fuente_encabezado
        celda.fill = relleno_obligatorio if requerido else relleno_opcional
        ws.column_dimensions[get_column_letter(col_idx)].width = 26

    ws.freeze_panes = 'A2'

    # --- Hoja de ayuda: qué significa cada columna + claves de carrera válidas ---
    ws_guia = wb.create_sheet('Guía')
    ws_guia.append(['Columna', 'Obligatoria', 'Descripción'])
    for celda in ws_guia[1]:
        celda.font = Font(bold=True)

    for nombre_col, requerido, desc in COLUMNAS_IMPORTACION_ALUMNOS:
        ws_guia.append([nombre_col, 'Sí' if requerido else 'No', desc])

    ws_guia.column_dimensions['A'].width = 28
    ws_guia.column_dimensions['B'].width = 12
    ws_guia.column_dimensions['C'].width = 60

    ws_guia.append([])
    fila_titulo_planes = ws_guia.max_row + 1
    ws_guia.cell(row=fila_titulo_planes, column=1, value='Claves de carrera disponibles:').font = Font(bold=True)

    for plan in PlanEstudio.query.filter_by(activo=True).order_by(PlanEstudio.clave_carrera.asc()).all():
        ws_guia.append([plan.clave_carrera, plan.nombre])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name='plantilla_carga_masiva_alumnos.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


@alumnos_bp.route('/alumnos/importar', methods=['GET', 'POST'])
@rol_requerido('DIRECTIVO')
def importar_alumnos():
    if request.method == 'GET':
        return render_template('importar_alumnos.html')

    archivo = request.files.get('archivo_excel')
    if not archivo or archivo.filename == '':
        flash('Selecciona un archivo Excel (.xlsx) para importar.', 'danger')
        return redirect(url_for('alumnos.importar_alumnos'))

    if not archivo.filename.lower().endswith('.xlsx'):
        flash('El archivo debe tener formato .xlsx (Excel). Usa la plantilla descargable.', 'danger')
        return redirect(url_for('alumnos.importar_alumnos'))

    error_archivo = error_de_tamano_xlsx(archivo.stream)
    if error_archivo:
        flash(error_archivo, 'danger')
        return redirect(url_for('alumnos.importar_alumnos'))

    try:
        wb = openpyxl.load_workbook(archivo, data_only=True)
        ws = wb['Alumnos'] if 'Alumnos' in wb.sheetnames else wb.active
    except Exception:
        flash('No se pudo leer el archivo. Verifica que sea un .xlsx válido generado con la plantilla.', 'danger')
        return redirect(url_for('alumnos.importar_alumnos'))

    error_hoja = error_de_dimensiones_hoja(ws)
    if error_hoja:
        flash(error_hoja, 'danger')
        return redirect(url_for('alumnos.importar_alumnos'))

    generar_cargos = request.form.get('generar_cargos') == 'on'

    encabezados = [(c.value.strip() if isinstance(c.value, str) else c.value) for c in ws[1]]
    nombres_columnas_conocidas = [nombre for nombre, _, _ in COLUMNAS_IMPORTACION_ALUMNOS]

    indice_columna = {
        nombre: encabezados.index(nombre)
        for nombre in nombres_columnas_conocidas
        if nombre in encabezados
    }

    faltantes = [
        nombre for nombre, requerido, _ in COLUMNAS_IMPORTACION_ALUMNOS
        if requerido and nombre not in indice_columna
    ]
    if faltantes:
        flash(
            f'Al archivo le faltan columnas obligatorias: {", ".join(faltantes)}. '
            'Descarga la plantilla más reciente y no cambies los nombres de las columnas.',
            'danger'
        )
        return redirect(url_for('alumnos.importar_alumnos'))

    def valor_de(fila, nombre_col):
        idx = indice_columna.get(nombre_col)
        if idx is None or idx >= len(fila):
            return None
        valor = fila[idx].value
        if isinstance(valor, str):
            valor = valor.strip()
        if valor in ('', None):
            return None
        if nombre_col in LONGITUDES_IMPORTACION and not isinstance(valor, str):
            valor = _texto_de_celda(valor)
        return valor

    planes_por_clave = {
        p.clave_carrera: p for p in PlanEstudio.query.filter_by(activo=True).all()
    }

    exitosos = []
    errores = []
    avisos_de_configuracion = []

    for num_fila, fila in enumerate(ws.iter_rows(min_row=2), start=2):
        if all(c.value in (None, '') for c in fila):
            continue  # Fila completamente vacía, se omite sin generar error

        fila_errores = []

        nombre_completo = valor_de(fila, 'nombre_completo')
        curp = str(valor_de(fila, 'curp') or '').upper()
        clave_carrera = str(valor_de(fila, 'clave_carrera') or '').upper()
        sexo = valor_de(fila, 'sexo')

        if not nombre_completo or len(str(nombre_completo)) < 5:
            fila_errores.append('nombre_completo inválido o vacío')
        elif str(nombre_completo)[0] in CARACTERES_FORMULA_EXCEL:
            fila_errores.append('nombre_completo empieza con =, +, - o @: Excel lo leería como fórmula al exportar')

        for columna, maximo in LONGITUDES_IMPORTACION.items():
            texto = valor_de(fila, columna)
            if texto is not None and len(str(texto)) > maximo:
                fila_errores.append(f'{columna} excede {maximo} caracteres')

        correo = valor_de(fila, 'correo')
        if correo and not CORREO_REGEX.match(str(correo)):
            fila_errores.append('correo con formato inválido')

        if not re.match(r'^[A-Z0-9]{18}$', curp):
            fila_errores.append('CURP inválida (deben ser 18 caracteres)')
        elif Alumno.query.filter_by(curp=curp).first():
            fila_errores.append(f'ya existe un alumno con la CURP {curp}')

        plan = planes_por_clave.get(clave_carrera)
        if not plan:
            fila_errores.append(f'clave_carrera "{clave_carrera}" no corresponde a ningún plan activo')

        if not sexo:
            fila_errores.append('sexo vacío')

        def parsear_fecha(nombre_columna):
            crudo = valor_de(fila, nombre_columna)
            if isinstance(crudo, datetime):
                return crudo.date(), None
            if isinstance(crudo, str):
                try:
                    return datetime.strptime(crudo, '%Y-%m-%d').date(), None
                except ValueError:
                    return None, f'{nombre_columna} con formato inválido (usa AAAA-MM-DD)'
            return None, f'{nombre_columna} vacía'

        fecha_nacimiento, error_fn = parsear_fecha('fecha_nacimiento')
        if error_fn:
            fila_errores.append(error_fn)

        fecha_certificado, error_fc = parsear_fecha('fecha_certificado_prepa')
        if error_fc:
            fila_errores.append(error_fc)

        estatus_raw = str(valor_de(fila, 'estatus') or 'ACTIVO').upper()
        if estatus_raw not in EstatusAlumno.__members__:
            fila_errores.append(f'estatus "{estatus_raw}" no es válido')

        turno = None
        turno_raw = valor_de(fila, 'turno')
        if turno_raw:
            turno_raw = str(turno_raw).upper()
            if turno_raw in TurnoAlumno.__members__:
                turno = TurnoAlumno[turno_raw]
            else:
                fila_errores.append(f'turno "{turno_raw}" no es válido')

        modalidad = None
        modalidad_raw = valor_de(fila, 'modalidad')
        if modalidad_raw:
            modalidad_raw = str(modalidad_raw).upper()
            if modalidad_raw in ModalidadEstudio.__members__:
                modalidad = ModalidadEstudio[modalidad_raw]
            else:
                fila_errores.append(f'modalidad "{modalidad_raw}" no es válida')

        cuatrimestre_raw = valor_de(fila, 'cuatrimestre_actual')
        cuatrimestre_actual = 1
        if cuatrimestre_raw is not None:
            try:
                cuatrimestre_actual = int(cuatrimestre_raw)
                if not (1 <= cuatrimestre_actual <= _max_periodos()):
                    fila_errores.append(f'cuatrimestre_actual debe estar entre 1 y {_max_periodos()}')
            except (TypeError, ValueError):
                fila_errores.append('cuatrimestre_actual debe ser un número')

        if fila_errores:
            errores.append({
                'fila': num_fila,
                'nombre': nombre_completo or '(sin nombre)',
                'errores': fila_errores,
            })
            continue

        try:
            alumno, error_creacion = crear_alumno_generando_matricula(
                plan,
                nombre_completo=nombre_completo,
                curp=curp,
                fecha_nacimiento=fecha_nacimiento,
                fecha_certificado_prepa=fecha_certificado,
                estatus=EstatusAlumno[estatus_raw],
                cuatrimestre_actual=cuatrimestre_actual,
                correo=valor_de(fila, 'correo'),
                telefono=valor_de(fila, 'telefono'),
                telefono_movil=valor_de(fila, 'telefono_movil'),
                sexo=sexo,
                numero_identificacion=valor_de(fila, 'numero_identificacion'),
                estado_civil=valor_de(fila, 'estado_civil'),
                nacionalidad=valor_de(fila, 'nacionalidad') or 'Mexicana',
                tipo_sangre=valor_de(fila, 'tipo_sangre'),
                domicilio_calle_numero=valor_de(fila, 'domicilio_calle_numero'),
                domicilio_ciudad=valor_de(fila, 'domicilio_ciudad'),
                domicilio_cp=valor_de(fila, 'domicilio_cp'),
                domicilio_estado=valor_de(fila, 'domicilio_estado'),
                contacto_emergencia_nombre=valor_de(fila, 'contacto_emergencia_nombre'),
                contacto_emergencia_telefono=valor_de(fila, 'contacto_emergencia_telefono'),
                contacto_emergencia_parentesco=valor_de(fila, 'contacto_emergencia_parentesco'),
                turno=turno,
                modalidad=modalidad,
                grupo_actual=valor_de(fila, 'grupo_actual'),
            )
        except SQLAlchemyError:
            # Un valor que la BD rechaza (DataError, etc.) solo descarta ESTA fila:
            # antes el 500 abortaba el lote y las filas siguientes se perdían sin aviso.
            db.session.rollback()
            errores.append({'fila': num_fila, 'nombre': nombre_completo, 'errores': ['la base de datos rechazó esta fila (revisa el formato de sus valores)']})
            continue

        if error_creacion:
            # commit por fila individual (ver crear_alumno_generando_matricula):
            # esta falla NO afecta a las filas anteriores, que ya quedaron
            # guardadas en la base de datos de forma independiente.
            errores.append({
                'fila': num_fila,
                'nombre': nombre_completo,
                'errores': [error_creacion],
            })
            continue

        # Alta completa (igual que la validación manual desde el expediente): un
        # alumno importado como no-Pendiente ya está validado, y el cambio queda en
        # el historial con quién lo importó. Sin esto, al reingresar desde una baja se
        # le cobraba "Inscripción" a alguien ya inscrito.
        estatus_importado = EstatusAlumno[estatus_raw]
        cargos_generados = []
        try:
            if estatus_importado != EstatusAlumno.PENDIENTE:
                alumno.fecha_validacion = ahora_utc()
            db.session.add(HistorialEstatus(
                matricula_fk=alumno.matricula_id, usuario_fk=current_user.id,
                estatus_anterior=None, estatus_nuevo=estatus_importado,
                comentario='Alta por importación masiva',
            ))
            if generar_cargos and estatus_importado == EstatusAlumno.ACTIVO:
                if cuatrimestre_actual == 1:
                    cargos_generados, avisos = _generar_cargos_de_inscripcion(alumno)
                else:
                    cargos_generados, avisos = _generar_cargos_de_reinscripcion(alumno)
                for aviso in avisos:
                    if aviso not in avisos_de_configuracion:
                        avisos_de_configuracion.append(aviso)
                _generar_carga_academica(alumno)
            db.session.commit()
        except SQLAlchemyError:
            db.session.rollback()
            errores.append({'fila': num_fila, 'nombre': nombre_completo, 'errores': [
                f'el alumno se creó ({alumno.matricula_id}) pero no se pudo completar su alta (historial/cargos); no lo vuelvas a importar, revísalo en su expediente'
            ]})
            continue

        exitosos.append({'fila': num_fila, 'nombre': nombre_completo, 'matricula': alumno.matricula_id, 'cargos': len(cargos_generados)})

    if exitosos or errores:
        registrar(
            'IMPORTACION_ALUMNOS', 'Alumno', None,
            f'{len(exitosos)} alumno(s) importados, {len(errores)} fila(s) con error; '
            f'cargos generados: {sum(e.get("cargos", 0) for e in exitosos)}; '
            f'generar_cargos={"sí" if generar_cargos else "no"}; archivo "{archivo.filename[:120]}"',
        )
        db.session.commit()

    if exitosos:
        flash(f'Se importaron {len(exitosos)} alumno(s) correctamente.', 'success')

    if errores:
        flash(f'{len(errores)} fila(s) no se pudieron importar (ver detalle abajo).', 'warning')
    for aviso in avisos_de_configuracion:
        flash(aviso, 'warning')

    return render_template('importar_alumnos.html', exitosos=exitosos, errores=errores)


@alumnos_bp.route('/alumno/<matricula>/expediente')
@rol_requerido('DIRECTIVO', 'ADMINISTRATIVO', 'CAPTURADOR')
def ver_expediente(matricula):
    alumno = db.get_or_404(Alumno, matricula)

    # Todas las materias del plan del alumno (el "Escudo": solo las de SU plan)
    materias_plan = (
        Materia.query
        .filter_by(id_plan_fk=alumno.id_plan_fk)
        .order_by(Materia.cuatrimestre.asc(), Materia.nombre.asc())
        .all()
    )

    # Última calificación registrada por materia (si una materia se recursó,
    # más de una Calificacion puede existir; nos quedamos con la más reciente)
    calif_por_materia = {}
    for c in sorted(alumno.calificaciones, key=lambda c: c.fecha_captura):
        calif_por_materia[c.id_materia_fk] = c

    # Agrupar por cuatrimestre para el acordeón
    cuatrimestres = {}
    for materia in materias_plan:
        cuatrimestres.setdefault(materia.cuatrimestre, []).append({
            'materia': materia,
            'calificacion': calif_por_materia.get(materia.id)
        })

    # Materias del plan que aún no están aprobadas (calificación >= 6).
    # Se usa para advertir (sin bloquear) al intentar marcar al alumno como Egresado.
    materias_no_aprobadas = [
        materia for materia in materias_plan
        if not (calif_por_materia.get(materia.id) and calif_por_materia[materia.id].calificacion_final >= 6)
    ]

    return render_template(
        'expediente.html',
        alumno=alumno,
        cuatrimestres=cuatrimestres,
        max_cuatrimestres=_max_periodos(),
        estatus_disponibles=list(EstatusAlumno),
        materias_no_aprobadas=materias_no_aprobadas,
        historial_estatus=alumno.historial_estatus
    )


@alumnos_bp.route('/alumno/<matricula>/ficha')
@rol_requerido('DIRECTIVO', 'ADMINISTRATIVO', 'CAPTURADOR')
def ficha_inscripcion(matricula):
    """
    Ficha de Inscripción imprimible (frente + reverso), a partir de los
    datos ya capturados en el registro público y el módulo de documentos.
    No requiere formulario propio: es una vista de solo lectura para
    imprimir/archivar, con checklist de documentos recibidos en el reverso.
    """
    alumno = db.get_or_404(Alumno, matricula)

    documentos_alumno = DocumentoAlumno.query.filter_by(matricula_fk=matricula).all()
    documentos_subidos = {doc.tipo_documento for doc in documentos_alumno}
    foto_alumno = next((doc for doc in documentos_alumno if doc.tipo_documento == TipoDocumento.FOTOGRAFIA), None)

    return render_template(
        'ficha_inscripcion.html',
        alumno=alumno,
        tipos_documento=list(TipoDocumento),
        documentos_subidos=documentos_subidos,
        foto_alumno=foto_alumno,
    )


@alumnos_bp.route('/alumno/<matricula>/cambiar-estatus', methods=['POST'])
@rol_requerido('DIRECTIVO', 'ADMINISTRATIVO', 'CAPTURADOR')
def cambiar_estatus(matricula):
    """
    Cambia el estatus del alumno (ej. Pendiente -> Activo -> Egresado).
    Registra el cambio en HistorialEstatus para auditoría. Si se intenta
    marcar como EGRESADO y el alumno tiene materias sin aprobar, se pide
    un comentario justificando el motivo (advertencia, no bloqueo total).
    """
    alumno = db.get_or_404(Alumno, matricula)
    nuevo_estatus_raw = request.form.get('nuevo_estatus', '')
    comentario = request.form.get('comentario', '').strip() or None

    if nuevo_estatus_raw not in EstatusAlumno.__members__:
        flash('El estatus indicado no es válido.', 'danger')
        return redirect(url_for('alumnos.ver_expediente', matricula=matricula))

    nuevo_estatus = EstatusAlumno[nuevo_estatus_raw]

    if nuevo_estatus == alumno.estatus:
        flash('El alumno ya tiene ese estatus; no se realizó ningún cambio.', 'warning')
        return redirect(url_for('alumnos.ver_expediente', matricula=matricula))

    if nuevo_estatus == EstatusAlumno.EGRESADO:
        materias_plan = Materia.query.filter_by(id_plan_fk=alumno.id_plan_fk).all()
        calif_por_materia = {c.id_materia_fk: c for c in alumno.calificaciones}
        faltantes = [
            m for m in materias_plan
            if not (calif_por_materia.get(m.id) and calif_por_materia[m.id].calificacion_final >= 6)
        ]

        if faltantes and not comentario:
            nombres = ', '.join(m.nombre for m in faltantes[:5])
            extra = f' y {len(faltantes) - 5} más' if len(faltantes) > 5 else ''
            flash(
                f'El alumno tiene {len(faltantes)} materia(s) sin aprobar ({nombres}{extra}). '
                'Si de verdad quieres marcarlo como Egresado, agrega un comentario '
                'explicando el motivo y vuelve a intentarlo.',
                'warning'
            )
            return redirect(url_for('alumnos.ver_expediente', matricula=matricula))

        if alumno.tiene_adeudo() and not comentario:
            detalle_adeudo = f' de ${alumno.saldo_total_adeudado():.2f}' if current_user.puede_cobros() else ''
            flash(
                f'El alumno tiene un adeudo económico{detalle_adeudo}. '
                'Si de verdad quieres marcarlo como Egresado, agrega un comentario '
                'explicando el motivo (ver Cobros para el detalle) y vuelve a intentarlo.',
                'warning'
            )
            return redirect(url_for('alumnos.ver_expediente', matricula=matricula))

    registro = HistorialEstatus(
        matricula_fk=alumno.matricula_id,
        usuario_fk=current_user.id,
        estatus_anterior=alumno.estatus,
        estatus_nuevo=nuevo_estatus,
        comentario=comentario,
    )
    alumno.estatus = nuevo_estatus

    cargos_auto_generados = []
    materias_auto_inscritas = []
    avisos_de_configuracion = []
    if nuevo_estatus == EstatusAlumno.ACTIVO and not alumno.fecha_validacion:
        alumno.fecha_validacion = ahora_utc()
        cargos_auto_generados, avisos_de_configuracion = _generar_cargos_de_inscripcion(alumno)
        materias_auto_inscritas = _generar_carga_academica(alumno)

    db.session.add(registro)
    if cargos_auto_generados:
        registrar(
            'CARGOS_AUTOMATICOS', 'Alumno', alumno.matricula_id,
            f'{len(cargos_auto_generados)} cargo(s) generados al activar: ' + ', '.join(f'{c.concepto} {c.periodo_escolar}' for c in cargos_auto_generados),
            matricula=alumno.matricula_id,
        )
    try:
        db.session.commit()
    except IntegrityError:
        # CONCURRENCIA: _cargo_duplicado() ya se revisó arriba (dentro de
        # _generar_cargos_de_inscripcion()), pero eso es solo un check en
        # Python -- la garantía real es el índice único de la BD (ver
        # migración b7e2c9a41f3d). Si dos peticiones activan al mismo
        # alumno casi al mismo tiempo, una gana y la otra cae aquí.
        db.session.rollback()
        flash(
            'No se pudo actualizar el estatus: ya se generó un cargo con la misma '
            'clave (alumno + concepto + periodo) justo ahora, probablemente por '
            'otra operación al mismo tiempo. Vuelve a intentarlo.',
            'danger'
        )
        return redirect(url_for('alumnos.ver_expediente', matricula=matricula))

    flash(f'Estatus del alumno actualizado a "{nuevo_estatus.value}".', 'success')
    if cargos_auto_generados:
        flash(
            f'Se generaron automáticamente {len(cargos_auto_generados)} cargo(s) de inscripción: '
            f'{", ".join(c.periodo_escolar for c in cargos_auto_generados)}. '
            'Revísalos en Cobros.',
            'success'
        )
    if materias_auto_inscritas:
        flash(
            f'Se generó su carga académica del {alumno.cuatrimestre_actual}° cuatrimestre '
            f'({len(materias_auto_inscritas)} materia(s)). Puedes verla en Carga Académica.',
            'success'
        )
    # Falta configuración de precios: no se inventa un monto, se dice qué falta.
    for aviso in avisos_de_configuracion:
        flash(aviso, 'warning')
    return redirect(url_for('alumnos.ver_expediente', matricula=matricula))


@alumnos_bp.route('/alumno/<matricula>/avanzar-cuatrimestre', methods=['POST'])
@rol_requerido('DIRECTIVO', 'ADMINISTRATIVO', 'CAPTURADOR')
def avanzar_cuatrimestre(matricula):
    """Avanza a UN alumno al siguiente cuatrimestre -- para casos sueltos (ej. alguien que regresó de Baja Temporal)."""
    alumno = db.get_or_404(Alumno, matricula)

    ok, mensaje, cargos, materias, avisos_de_configuracion = _avanzar_cuatrimestre(alumno)

    if ok:
        if cargos:
            registrar(
                'CARGOS_AUTOMATICOS', 'Alumno', alumno.matricula_id,
                f'{len(cargos)} cargo(s) generados al avanzar al {alumno.cuatrimestre_actual}° cuatrimestre',
                matricula=alumno.matricula_id,
            )
        try:
            db.session.commit()
        except IntegrityError:
            # CONCURRENCIA: mismo caso que cambiar_estatus() -- ver migración
            # b7e2c9a41f3d.
            db.session.rollback()
            flash(
                'No se pudo avanzar de cuatrimestre: ya se generó un cargo con la '
                'misma clave (alumno + concepto + periodo) justo ahora, '
                'probablemente por otra operación al mismo tiempo. Vuelve a '
                'intentarlo.',
                'danger'
            )
            return redirect(url_for('alumnos.ver_expediente', matricula=matricula))
        flash(mensaje, 'success')
        if cargos:
            flash(f'Se generaron {len(cargos)} cargo(s) nuevo(s). Revísalos en Cobros.', 'success')
        if materias:
            flash(f'Se generó su carga académica del {alumno.cuatrimestre_actual}° cuatrimestre ({len(materias)} materia(s)).', 'success')
        for aviso in avisos_de_configuracion:
            flash(aviso, 'warning')
    else:
        db.session.rollback()
        flash(mensaje, 'danger')

    return redirect(url_for('alumnos.ver_expediente', matricula=matricula))


@alumnos_bp.route('/alumnos/avanzar-cuatrimestre-lote', methods=['GET', 'POST'])
@rol_requerido('DIRECTIVO', 'ADMINISTRATIVO')
def avanzar_cuatrimestre_lote():
    """
    Avanza de golpe a TODOS los alumnos Activos de una carrera+cuatrimestre
    al siguiente cuatrimestre -- pensado para el cambio de periodo, cuando
    todo un grupo pasa junto. Reutiliza _avanzar_cuatrimestre(), el mismo
    helper que la versión individual, así que genera exactamente los
    mismos cargos y carga académica, solo que para muchos a la vez.
    """
    planes = PlanEstudio.query.filter_by(activo=True).order_by(PlanEstudio.nombre.asc()).all()
    max_cuatri = _max_periodos()

    if request.method == 'POST':
        plan_id = request.form.get('plan_id', type=int)
        cuatrimestre_actual = request.form.get('cuatrimestre_actual', type=int)

        plan = db.session.get(PlanEstudio, plan_id) if plan_id else None
        if not plan or not cuatrimestre_actual:
            flash('Selecciona carrera y el cuatrimestre en el que están ahora.', 'danger')
            return render_template('avanzar_cuatrimestre.html', planes=planes, max_cuatrimestres=max_cuatri)

        alumnos = (
            Alumno.query
            .filter_by(id_plan_fk=plan.id, cuatrimestre_actual=cuatrimestre_actual, estatus=EstatusAlumno.ACTIVO)
            .order_by(Alumno.nombre_completo.asc())
            .all()
        )

        avanzados = []
        omitidos = []
        # Configuración de precios que falta. Se junta SIN repetir: si a 40
        # alumnos les falta el mismo precio, es un solo problema de
        # configuración, no 40 avisos en pantalla.
        avisos_de_configuracion = []
        # no_autoflush: _avanzar_cuatrimestre() consulta la BD (conceptos,
        # _cargo_duplicado) por cada alumno del lote -- sin esto, esa
        # consulta dispara un autoflush de los INSERT de alumnos ya
        # procesados en este mismo lote, y si alguno choca con el índice
        # único, el IntegrityError se dispara AQUÍ (fuera del try/except
        # de abajo) en vez de en el commit explícito, y el lote entero
        # truena con un 500 en lugar del mensaje de abajo.
        with db.session.no_autoflush:
            for alumno in alumnos:
                ok, mensaje, cargos, materias, avisos = _avanzar_cuatrimestre(alumno)
                if ok:
                    avanzados.append({'alumno': alumno, 'cargos': len(cargos), 'materias': len(materias)})
                else:
                    omitidos.append({'alumno': alumno, 'motivo': mensaje})
                for aviso in avisos:
                    if aviso not in avisos_de_configuracion:
                        avisos_de_configuracion.append(aviso)

        registrar(
            'AVANCE_LOTE', 'PlanEstudio', plan.id,
            f'{plan.nombre}, del {cuatrimestre_actual}° al {cuatrimestre_actual + 1}°: {len(avanzados)} avanzaron '
            f'({sum(a["cargos"] for a in avanzados)} cargos generados), {len(omitidos)} omitidos',
        )
        try:
            db.session.commit()
        except IntegrityError:
            # CONCURRENCIA: mismo caso que generar_mensualidades() -- todo el
            # lote es una sola transacción; si algún cargo choca con el
            # índice único (migración b7e2c9a41f3d) se descarta el lote
            # COMPLETO (nunca a medias) y se pide reintentar.
            db.session.rollback()
            flash(
                'No se aplicó el avance de cuatrimestre: alguno de estos cargos '
                'ya se generó justo ahora por otra operación. Vuelve a '
                'intentarlo -- los que ya estén al día se omitirán solos.',
                'danger'
            )
            return render_template('avanzar_cuatrimestre.html', planes=planes, max_cuatrimestres=max_cuatri)

        flash(
            f'{len(avanzados)} alumno(s) avanzaron al {cuatrimestre_actual + 1}° cuatrimestre. '
            f'{len(omitidos)} se omitieron (ver detalle abajo).',
            'success' if avanzados else 'warning'
        )
        for aviso in avisos_de_configuracion:
            flash(aviso, 'warning')
        return render_template(
            'avanzar_cuatrimestre.html',
            planes=planes,
            max_cuatrimestres=max_cuatri,
            avanzados=avanzados,
            omitidos=omitidos,
        )

    return render_template('avanzar_cuatrimestre.html', planes=planes, max_cuatrimestres=max_cuatri)

