"""
Cada institución tiene su propio formato de matrícula y, al migrar, sus alumnos YA traen matrícula.
Antes el formato estaba fijo en el código (LEN2026-00001) y la importación siempre inventaba una.
Ahora: formato configurable (solo para matrículas NUEVAS) y matrícula propia opcional al importar.
"""

import io

import pytest
from openpyxl import Workbook

from tests.conftest import crear_plan, crear_alumno, crear_usuario, login
from app import db, Alumno, ConfiguracionInstitucion, ahora_utc, PlanEstudio
from modelos import BitacoraAuditoria
from servicios.matriculas import generar_matricula, ejemplo_matricula

ANIO = ahora_utc().year


def _formato(**cambios):
    cfg = ConfiguracionInstitucion.obtener()
    for k, v in cambios.items():
        setattr(cfg, k, v)
    db.session.commit()
    return cfg


def _directivo(client):
    crear_usuario(username='dir1'); login(client, 'dir1', 'clave12345')


# ------------------------------- generación -------------------------------

def test_el_formato_por_defecto_sigue_siendo_el_de_siempre(app):
    plan = crear_plan(clave='LEN')
    assert generar_matricula(plan) == f'LEN{ANIO}-00001'


def test_prefijo_fijo_sin_clave_con_anio_y_separador_vacio(app):
    _formato(matricula_prefijo='UNI', matricula_incluye_clave=False, matricula_incluye_anio=True, matricula_separador='', matricula_digitos=6)
    a, b = crear_plan(clave='LEN', anio=2026), crear_plan(nombre='Otra', clave='ISC', anio=2026)
    assert generar_matricula(a) == f'UNI{ANIO}000001'
    crear_alumno(a, matricula_id=f'UNI{ANIO}000001')
    assert generar_matricula(b) == f'UNI{ANIO}000002'          # sin clave en el formato, las carreras comparten consecutivo


def test_solo_consecutivo(app):
    _formato(matricula_prefijo='', matricula_incluye_clave=False, matricula_incluye_anio=False, matricula_separador='', matricula_digitos=4)
    plan = crear_plan()
    assert generar_matricula(plan) == '0001'
    crear_alumno(plan, matricula_id='0001')
    assert generar_matricula(plan) == '0002'


def test_el_consecutivo_continua_despues_de_matriculas_traidas_de_otro_sistema(app):
    plan = crear_plan(clave='LEN')
    crear_alumno(plan, curp='AAAA010101HDFXYZ01', matricula_id=f'LEN{ANIO}-00050')     # numérica con el mismo prefijo
    crear_alumno(plan, curp='BBBB010101HDFXYZ02', matricula_id=f'LEN{ANIO}-ABC')       # con el prefijo pero no numérica: se ignora
    crear_alumno(plan, curp='CCCC010101HDFXYZ03', matricula_id='MAT-9999')             # formato ajeno
    assert generar_matricula(plan) == f'LEN{ANIO}-00051'


def test_el_ejemplo_refleja_la_configuracion(app):
    cfg = _formato(matricula_prefijo='ENF', matricula_separador='.')
    assert ejemplo_matricula(cfg, 'LEN', 2026) == 'ENFLEN2026.00001'


# ------------------------------- configuración desde la pantalla -------------------------------

def _post_config(client, **cambios):
    datos = dict(nombre_institucion='Escuela de Prueba', nombre_periodo_singular='Cuatrimestre', nombre_periodo_plural='Cuatrimestres',
                 nombre_programa_singular='Carrera', nombre_programa_plural='Carreras', max_periodos='9',
                 matricula_prefijo='', matricula_incluye_clave='on', matricula_incluye_anio='on', matricula_separador='-', matricula_digitos='5')
    datos.update(cambios)
    return client.post('/configuracion/institucion', data={k: v for k, v in datos.items() if v is not None})


def test_guardar_el_formato_de_matricula_desde_la_pantalla_y_dejar_bitacora(client, app):
    crear_plan(clave='LEN'); _directivo(client)
    _post_config(client, matricula_prefijo='UNI', matricula_incluye_clave=None, matricula_separador='.', matricula_digitos='6')
    db.session.expire_all()
    cfg = ConfiguracionInstitucion.obtener()
    assert (cfg.matricula_prefijo, cfg.matricula_incluye_clave, cfg.matricula_incluye_anio, cfg.matricula_separador, cfg.matricula_digitos) == ('UNI', False, True, '.', 6)
    (reg,) = BitacoraAuditoria.query.filter_by(accion='CONFIG_MATRICULA').all()
    assert 'UNI' in reg.detalle
    assert f'UNI{ANIO}.000001' in client.get('/configuracion/institucion').get_data(as_text=True)     # ejemplo en pantalla


@pytest.mark.parametrize('campo,valor', [
    ('matricula_digitos', '2'), ('matricula_digitos', '9'), ('matricula_digitos', 'x'), ('matricula_digitos', '²'),
    ('matricula_separador', '/'), ('matricula_separador', 'abc'), ('matricula_separador', ' '),
    ('matricula_prefijo', 'ABC DEF'), ('matricula_prefijo', 'ABCDEFG'), ('matricula_prefijo', 'A/B'), ('matricula_prefijo', '%'),
])
def test_el_formato_de_matricula_rechaza_valores_invalidos(client, app, campo, valor):
    crear_plan(clave='LEN'); _directivo(client)
    r = _post_config(client, **{campo: valor})
    assert r.status_code == 302
    db.session.expire_all()
    cfg = ConfiguracionInstitucion.obtener()
    assert (cfg.matricula_prefijo, cfg.matricula_separador, cfg.matricula_digitos) == ('', '-', 5)


def test_no_se_acepta_un_formato_que_haga_matriculas_de_mas_de_20_caracteres(client, app):
    crear_plan(nombre='Clave larga', clave='ABCDEFGHIJ')            # 10 caracteres
    _directivo(client)
    _post_config(client, matricula_prefijo='UNIVER', matricula_digitos='8')      # 6+10+4+1+8 = 29 > 20
    db.session.expire_all()
    assert ConfiguracionInstitucion.obtener().matricula_prefijo == ''


def test_cambiar_el_formato_no_toca_las_matriculas_ya_emitidas(client, app):
    plan = crear_plan(clave='LEN'); a = crear_alumno(plan, matricula_id='LEN2026-00001')
    _directivo(client)
    _post_config(client, matricula_prefijo='UNI')
    db.session.expire_all()
    assert db.session.get(Alumno, 'LEN2026-00001') is not None


def test_el_registro_publico_usa_el_formato_configurado(client, app):
    plan = crear_plan(clave='LEN'); _formato(matricula_prefijo='ENF', matricula_incluye_clave=False, matricula_separador='.', matricula_digitos=4)
    client.post('/registro', data=dict(
        nombre_completo='Maria Fernanda Lopez', curp='LOPM050101MDFPRR09', fecha_nacimiento='2005-01-01', fecha_certificado_prepa='2023-07-01',
        id_plan_fk=plan.id, sexo='Femenino', domicilio_calle_numero='Calle 1', domicilio_ciudad='CDMX', domicilio_cp='01000', domicilio_estado='CDMX',
        contacto_emergencia_nombre='Mama', contacto_emergencia_telefono='5500000000', turno='MATUTINO', modalidad='ESCOLARIZADO'))
    assert Alumno.query.one().matricula_id == f'ENF{ANIO}.0001'


# ------------------------------- importación con matrícula propia -------------------------------

COLS = ['matricula', 'nombre_completo', 'curp', 'fecha_nacimiento', 'fecha_certificado_prepa', 'clave_carrera', 'sexo']


def _xlsx(filas, columnas=COLS):
    wb = Workbook(); ws = wb.active; ws.append(columnas)
    for f in filas:
        ws.append(f)
    b = io.BytesIO(); wb.save(b); b.seek(0); return b


def _fila(matricula, i, nombre=None):
    return [matricula, nombre or f'Alumno Numero {i}', f'MASI{i:014d}', '2005-01-01', '2023-07-01', 'LEN', 'Femenino']


def _importar(client, filas, columnas=COLS):
    return client.post('/alumnos/importar', data={'archivo_excel': (_xlsx(filas, columnas), 'a.xlsx')}, content_type='multipart/form-data', follow_redirects=True)


def test_la_importacion_respeta_la_matricula_que_ya_trae_cada_alumno(client, app):
    crear_plan(clave='LEN'); _directivo(client)
    _importar(client, [_fila('UNI-0001', 1), _fila(2024001234, 2), _fila(None, 3)])
    assert {a.matricula_id for a in Alumno.query.all()} == {'UNI-0001', '2024001234', f'LEN{ANIO}-00001'}


@pytest.mark.parametrize('matricula', ['AB CD', '../x', 'a/b', 'X' * 21, '-empieza', 'ñandú1', 'a;b'])
def test_la_importacion_rechaza_matriculas_con_caracteres_o_largo_invalidos(client, app, matricula):
    crear_plan(clave='LEN'); _directivo(client)
    r = _importar(client, [_fila(matricula, 1)])
    assert Alumno.query.count() == 0
    assert 'matricula' in r.get_data(as_text=True).lower()


def test_la_importacion_rechaza_una_matricula_repetida_en_la_base_o_en_el_mismo_archivo(client, app):
    plan = crear_plan(clave='LEN'); crear_alumno(plan, matricula_id='UNI-0001'); _directivo(client)
    r = _importar(client, [_fila('UNI-0001', 1), _fila('UNI-0002', 2), _fila('UNI-0002', 3)])
    assert sorted(a.matricula_id for a in Alumno.query.all()) == ['UNI-0001', 'UNI-0002']
    assert Alumno.query.filter_by(matricula_id='UNI-0002').count() == 1
    assert r.get_data(as_text=True).count('ya existe') >= 2


def test_sin_columna_de_matricula_se_generan_automaticas_como_siempre(client, app):
    crear_plan(clave='LEN'); _directivo(client)
    cols = COLS[1:]
    _importar(client, [f[1:] for f in (_fila('x', 1), _fila('x', 2))], columnas=cols)
    assert sorted(a.matricula_id for a in Alumno.query.all()) == [f'LEN{ANIO}-00001', f'LEN{ANIO}-00002']


def test_la_plantilla_de_importacion_trae_la_columna_de_matricula_como_opcional(client, app):
    import openpyxl
    crear_plan(clave='LEN'); _directivo(client)
    wb = openpyxl.load_workbook(io.BytesIO(client.get('/alumnos/importar/plantilla').data))
    assert 'matricula' in [c.value for c in wb['Alumnos'][1]]
    guia = {fila[0]: fila[1] for fila in wb['Guía'].iter_rows(values_only=True) if fila[0]}
    assert guia['matricula'] == 'No'
