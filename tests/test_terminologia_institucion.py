"""
Configuración de la Institución: cada institución define su nombre, cómo llama a sus periodos y a su programa,
cuántos periodos tiene, el género gramatical del programa y el formato de matrícula. Revisión 2026-09-20:
los cambios SE GUARDABAN, pero el sistema no los respetaba: el nombre no se usaba en ninguna pantalla
(7 documentos decían "Universidad de Enfermería" escrito a mano) y ~35 mensajes seguían diciendo "cuatrimestre"
y "carrera". Estas pruebas recorren cada concepto de punta a punta.
"""

import io
import re
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from tests.conftest import crear_plan, crear_materia, crear_alumno, crear_usuario, login
from app import (
    db, mail, Alumno, Calificacion, Cargo, ConceptoCobro, ConfiguracionInstitucion, EstatusAlumno,
    InscripcionMateria, Pago, RolUsuario,
)
from modelos import BitacoraAuditoria

NUEVOS = dict(
    nombre_institucion='Instituto QA-7 del Valle', nombre_periodo_singular='Trimestre', nombre_periodo_plural='Trimestres',
    nombre_programa_singular='Licenciatura', nombre_programa_plural='Licenciaturas', max_periodos='7', programa_genero='femenino',
)
PALABRAS_VIEJAS = re.compile(r'(?i)(universidad de enfermer|cuatrimestre|carrera)')


def _visible(html):
    html = re.sub(r'(?is)<(script|style)[^>]*>.*?</\1>', ' ', html)
    atributos = ' '.join(re.findall(r'(?:placeholder|title|aria-label|data-[a-z]+)="([^"]*)"', html))
    return re.sub(r'\s+', ' ', re.sub(r'(?s)<[^>]+>', ' ', html) + ' ' + atributos)


def _anonimo_get(client, url):
    """Petición sin sesión: contexto propio (la fixture mantiene un app_context y Flask-Login cachea current_user en `g`)."""
    with client.application.app_context():
        return client.application.test_client().get(url)


def _guardar(client, **cambios):
    datos = dict(NUEVOS); datos.update(cambios)
    return client.post('/configuracion/institucion', data={k: v for k, v in datos.items() if v is not None})


def _config():
    db.session.expire_all()
    return ConfiguracionInstitucion.obtener()


def _mensajes(client):
    with client.session_transaction() as s:
        return ' | '.join(m for _, m in s.pop('_flashes', []))


@pytest.fixture
def escenario(client, app):
    plan = crear_plan(nombre='Plan Demo Uno', clave='PDU')
    plan.monto_mensualidad = Decimal('2500'); db.session.commit()
    materia = crear_materia(plan, 'Materia Demo', 1)
    alumno = crear_alumno(plan, matricula_id='PDU2026-00001', nombre='Alumno Demo Uno', curp='DEMO010101HDFXYZ01')
    alumno.correo = 'alumno@ejemplo.mx'
    db.session.add(Calificacion(matricula_fk=alumno.matricula_id, id_materia_fk=materia.id, calificacion_final=9.0, periodo_escolar='2026-A'))
    db.session.add(InscripcionMateria(matricula_fk=alumno.matricula_id, id_materia_fk=materia.id, periodo_escolar='2026-A'))
    conc = ConceptoCobro(nombre='Mensualidad', es_mensualidad=True, activo=True); db.session.add(conc); db.session.commit()
    cargo = Cargo(matricula_fk=alumno.matricula_id, concepto_cobro_fk=conc.id, concepto='Mensualidad', monto=Decimal('2500'), periodo_escolar='2026-C-Sep')
    db.session.add(cargo); db.session.commit()
    pago = Pago(cargo_fk=cargo.id, monto_pagado=Decimal('500'), folio='PAGO-2026-000001'); db.session.add(pago); db.session.commit()
    crear_usuario(username='dir1'); login(client, 'dir1', 'clave12345')
    return dict(plan=plan, materia=materia, alumno=alumno, cargo=cargo, pago=pago)


def _paginas(e):
    m = e['alumno'].matricula_id
    return ['/', f'/alumno/{m}/expediente', f'/alumno/{m}/ficha', f'/alumno/{m}/boleta', f'/alumno/{m}/carga-academica',
            f'/alumno/{m}/historial-calificaciones', f'/alumno/{m}/cobros', f'/alumno/{m}/estado-cuenta', f'/alumno/{m}/becas',
            f'/alumno/{m}/documentos', f"/pago/{e['pago'].id}/recibo", '/planes/mensualidades', f"/planes/{e['plan'].id}/materias",
            '/conceptos-cobro', '/configuracion/cobros', '/cobros/dashboard', '/cobros/generar-mensualidades',
            '/cobros/recordatorios-vencimiento', '/reportes/cartera-vencida', '/reportes/cobros-del-dia', '/alumnos/importar',
            '/alumnos/avanzar-cuatrimestre-lote', '/boletas/importar', '/usuarios', '/usuarios/nuevo', '/perfil', '/registro']   # /auditoria a propósito NO: guarda el historial con los nombres anteriores


# ------------------------------- 1. nombre de la institución -------------------------------

def test_el_nombre_de_la_institucion_se_guarda(client, escenario):
    _guardar(client)
    assert _config().nombre_institucion == 'Instituto QA-7 del Valle'
    assert 'Instituto QA-7 del Valle' in client.get('/configuracion/institucion').get_data(as_text=True)


@pytest.mark.parametrize('ruta', ['ficha', 'boleta', 'carga-academica', 'estado-cuenta'])
def test_los_documentos_del_alumno_llevan_el_nombre_configurado(client, escenario, ruta):
    _guardar(client)
    html = _visible(client.get(f"/alumno/{escenario['alumno'].matricula_id}/{ruta}").get_data(as_text=True))
    assert 'Instituto QA-7 del Valle'.lower() in html.lower()
    assert 'universidad de enfermer' not in html.lower()


def test_el_recibo_de_pago_lleva_el_nombre_configurado(client, escenario):
    _guardar(client)
    html = _visible(client.get(f"/pago/{escenario['pago'].id}/recibo").get_data(as_text=True))
    assert 'Instituto QA-7 del Valle' in html and 'universidad de enfermer' not in html.lower()


def test_el_registro_publico_y_el_login_llevan_el_nombre_configurado(client, escenario):
    _guardar(client)
    assert 'Instituto QA-7 del Valle' in _visible(_anonimo_get(client, '/registro').get_data(as_text=True))
    assert 'Instituto QA-7 del Valle' in _visible(_anonimo_get(client, '/login').get_data(as_text=True))


def test_un_nombre_con_html_no_se_ejecuta_en_ninguna_pantalla(client, escenario):
    _guardar(client, nombre_institucion='<script>alert(1)</script>Uni')
    for url in (f"/alumno/{escenario['alumno'].matricula_id}/ficha", '/', f"/pago/{escenario['pago'].id}/recibo"):
        html = client.get(url).get_data(as_text=True)
        assert '<script>alert(1)</script>' not in html
    assert '&lt;script&gt;alert(1)&lt;/script&gt;Uni' in client.get(f"/alumno/{escenario['alumno'].matricula_id}/ficha").get_data(as_text=True)   # se muestra como TEXTO
    for url in ('/registro', '/login'):
        assert '<script>alert(1)</script>' not in _anonimo_get(client, url).get_data(as_text=True)


@pytest.mark.parametrize('nombre', ['ab', '   ', 'N' * 151, 'Instituto\nSaltos', 'Instituto\x07Campana'])
def test_el_nombre_de_la_institucion_valida_largo_y_caracteres_de_control(client, escenario, nombre):
    antes = _config().nombre_institucion
    r = _guardar(client, nombre_institucion=nombre)
    assert r.status_code == 302
    assert _config().nombre_institucion == antes


def test_el_nombre_se_limpia_de_espacios_sobrantes(client, escenario):
    _guardar(client, nombre_institucion='   Instituto   QA-7   del   Valle  ')
    assert _config().nombre_institucion == 'Instituto QA-7 del Valle'


def test_sin_nombre_configurado_el_director_ve_un_aviso_para_ponerlo(client, escenario):
    assert 'Configura el nombre de tu institución' in client.get('/').get_data(as_text=True)
    _guardar(client)
    assert 'Configura el nombre de tu institución' not in client.get('/').get_data(as_text=True)


# ------------------------------- 2. periodo y programa -------------------------------

def test_ninguna_pantalla_conserva_las_palabras_viejas_tras_cambiar_la_terminologia(client, escenario):
    _guardar(client)
    culpables = []
    for url in _paginas(escenario):
        r = client.get(url)
        assert r.status_code == 200, url
        for m in PALABRAS_VIEJAS.finditer(_visible(r.get_data(as_text=True))):
            culpables.append(f'{url}: ...{_visible(r.get_data(as_text=True))[max(0, m.start() - 25):m.end() + 25]}...')
            break
    assert not culpables, 'Palabras viejas visibles:\n' + '\n'.join(culpables)


@pytest.mark.parametrize('campo,valor', [
    ('nombre_periodo_singular', ''), ('nombre_periodo_plural', ''), ('nombre_programa_singular', ''), ('nombre_programa_plural', ''),
    ('nombre_periodo_singular', 'T' * 41), ('nombre_programa_plural', 'L' * 41), ('nombre_programa_singular', 'Lic\nenciatura'),
    ('nombre_periodo_plural', '<b>'),
])
def test_los_nombres_de_periodo_y_programa_se_validan(client, escenario, campo, valor):
    antes = _config().nombre_periodo_singular, _config().nombre_periodo_plural, _config().nombre_programa_singular, _config().nombre_programa_plural
    assert _guardar(client, **{campo: valor}).status_code == 302
    c = _config()
    assert (c.nombre_periodo_singular, c.nombre_periodo_plural, c.nombre_programa_singular, c.nombre_programa_plural) == antes


def test_los_mensajes_del_sistema_usan_la_terminologia_configurada(client, escenario):
    e = escenario; m = e['alumno'].matricula_id
    _guardar(client); _mensajes(client)
    todos = []

    client.post(f'/alumno/{m}/boleta', data={'periodo_escolar': '2026-B', 'cuatrimestre': '1', f"calificacion_{e['materia'].id}": '8'}); todos.append(_mensajes(client))
    client.get('/boletas/importar/plantilla'); todos.append(_mensajes(client))
    client.get(f"/boletas/importar/plantilla?plan_id={e['plan'].id}&cuatrimestre=3"); todos.append(_mensajes(client))
    client.post('/boletas/importar', data={'periodo_escolar': '2026-B'}); todos.append(_mensajes(client))
    client.post('/alumnos/avanzar-cuatrimestre-lote', data={}); todos.append(_mensajes(client))
    client.post(f'/alumno/{m}/avanzar-cuatrimestre'); todos.append(_mensajes(client))
    client.post(f'/alumno/{m}/documentos', data={'cuatrimestre_actual': '99'}); todos.append(_mensajes(client))
    client.post(f"/planes/{e['plan'].id}/materias", data={'nombre': 'Materia Demo', 'cuatrimestre': '1'}); todos.append(_mensajes(client))
    client.post(f"/planes/{e['plan'].id}/materias", data={'nombre': 'Otra Materia', 'cuatrimestre': '8'}); todos.append(_mensajes(client))
    client.post(f"/planes/{e['plan'].id}/materias", data={'nombre': 'Materia Nueva', 'cuatrimestre': '2'}); todos.append(_mensajes(client))
    client.post('/planes/nuevo', data=dict(nombre='Programa Nuevo', clave_carrera='NUE', anio_generacion='2026')); todos.append(_mensajes(client))
    client.post('/planes/nuevo', data=dict(nombre='Programa Nuevo', clave_carrera='NUE', anio_generacion='2026')); todos.append(_mensajes(client))
    client.post(f"/planes/{e['plan'].id}/toggle"); todos.append(_mensajes(client))
    anonimo = client.application.test_client()
    anonimo.post('/registro', data={'nombre_completo': 'x'})
    texto = ' | '.join(todos)
    assert not PALABRAS_VIEJAS.search(texto.replace('cuatrimestre_actual', '')), f'Mensajes con palabras viejas: {texto}'
    assert 'Trimestre'.lower() in texto.lower() and 'Licenciatura'.lower() in texto.lower()


def test_la_activacion_y_el_avance_de_un_alumno_hablan_con_la_terminologia_configurada(client, escenario):
    e = escenario
    e['plan'].monto_mensualidad = Decimal('2500'); db.session.commit()
    nuevo = crear_alumno(e['plan'], estatus=EstatusAlumno.PENDIENTE, matricula_id='PDU2026-00002', curp='DEMO020202HDFXYZ02')
    db.session.add(ConceptoCobro(nombre='Inscripción', monto_sugerido=Decimal('3000'), activo=True)); db.session.commit()
    _guardar(client); _mensajes(client)
    client.post(f'/alumno/{nuevo.matricula_id}/cambiar-estatus', data={'nuevo_estatus': 'ACTIVO'}); m1 = _mensajes(client)
    client.post(f'/alumno/{nuevo.matricula_id}/avanzar-cuatrimestre'); m2 = _mensajes(client)
    assert not PALABRAS_VIEJAS.search(m1 + m2), m1 + m2
    assert 'Trimestre'.lower() in (m1 + m2).lower()


def test_las_plantillas_de_excel_usan_la_terminologia_configurada(client, escenario):
    e = escenario; _guardar(client)
    guia = openpyxl.load_workbook(io.BytesIO(client.get('/alumnos/importar/plantilla').data))['Guía']
    texto = ' '.join(str(c) for fila in guia.iter_rows(values_only=True) for c in fila if c)
    assert 'Licenciatura' in texto or 'licenciatura' in texto
    assert not PALABRAS_VIEJAS.search(texto.replace('cuatrimestre_actual', '').replace('clave_carrera', '')), texto
    guia2 = openpyxl.load_workbook(io.BytesIO(client.get(f"/boletas/importar/plantilla?plan_id={e['plan'].id}&cuatrimestre=1").data))['Guía']
    filas = [fila[0] for fila in guia2.iter_rows(values_only=True) if fila[0]]
    assert 'Licenciatura' in filas and 'Trimestre' in filas


# ------------------------------- 3. género gramatical del programa -------------------------------

def test_el_genero_del_programa_se_guarda_y_se_refleja_en_los_textos(client, escenario):
    e = escenario
    vacio = crear_plan(nombre='Plan Sin Materias', clave='VAC')
    _guardar(client, nombre_programa_singular='Programa', nombre_programa_plural='Programas', programa_genero='masculino')
    assert _config().programa_es_femenino is False
    html = _visible(client.get(f"/planes/{vacio.id}/materias").get_data(as_text=True))
    assert 'Este programa' in html
    client.post('/planes/nuevo', data=dict(nombre='Otro Programa', clave_carrera='OTR', anio_generacion='2026')); client.post('/planes/nuevo', data=dict(nombre='Otro Programa', clave_carrera='OTR', anio_generacion='2026'))
    assert 'un programa' in _mensajes(client).lower()
    _guardar(client, nombre_programa_singular='Carrera', nombre_programa_plural='Carreras', programa_genero='femenino')
    assert _config().programa_es_femenino is True
    assert 'Esta carrera' in _visible(client.get(f"/planes/{vacio.id}/materias").get_data(as_text=True))


def test_un_genero_invalido_se_rechaza(client, escenario):
    _guardar(client, programa_genero='neutro')
    assert _config().nombre_periodo_singular == 'Cuatrimestre'


# ------------------------------- 4. máximo de periodos -------------------------------

def test_el_maximo_de_periodos_gobierna_materias_avance_y_selectores(client, escenario):
    e = escenario
    _guardar(client, max_periodos='3')
    assert _config().max_periodos == 3
    client.post(f"/planes/{e['plan'].id}/materias", data={'nombre': 'Materia Cuatro', 'cuatrimestre': '4'}); assert 'Materia Cuatro' not in [m.nombre for m in e['plan'].materias]
    client.post(f"/planes/{e['plan'].id}/materias", data={'nombre': 'Materia Tres', 'cuatrimestre': '3'}); assert 'Materia Tres' in [m.nombre for m in e['plan'].materias]
    boleta = client.get(f"/alumno/{e['alumno'].matricula_id}/boleta").get_data(as_text=True)
    assert re.findall(r'<option value="(\d+)"', boleta.split('name="cuatrimestre"')[1].split('</select>')[0]) == ['1', '2', '3']
    e['alumno'].cuatrimestre_actual = 3; db.session.commit()
    client.post(f"/alumno/{e['alumno'].matricula_id}/avanzar-cuatrimestre")
    db.session.expire_all(); assert db.session.get(Alumno, e['alumno'].matricula_id).cuatrimestre_actual == 3     # ya está en el máximo


@pytest.mark.parametrize('valor', ['0', '31', 'x', '', '²', '-1', '3.5'])
def test_el_maximo_de_periodos_valida_su_rango(client, escenario, valor):
    assert _guardar(client, max_periodos=valor).status_code == 302
    assert _config().max_periodos == 9


def test_no_se_puede_bajar_el_maximo_por_debajo_de_lo_que_ya_se_usa(client, escenario):
    """Antes: bajar de 9 a 7 con materias en el 8° las hacía DESAPARECER de la pantalla de materias (sin poder verlas ni editarlas)."""
    e = escenario
    crear_materia(e['plan'], 'Materia Ocho', 8)
    _guardar(client, max_periodos='7')
    assert _config().max_periodos == 9
    assert 'materias o alumnos en 8°' in _mensajes(client)      # explica por qué no se puede bajar el máximo
    e['alumno'].cuatrimestre_actual = 9; db.session.commit()
    _guardar(client, max_periodos='9'); assert _config().max_periodos == 9
    Materia = type(e['materia']); Materia.query.filter_by(nombre='Materia Ocho').delete(); db.session.commit()
    _guardar(client, max_periodos='8'); assert _config().max_periodos == 9        # el alumno sigue en el 9°


# ------------------------------- 5. bitácora, permisos, correos, exportaciones -------------------------------

def test_guardar_la_configuracion_queda_en_la_bitacora_con_antes_y_despues(client, escenario):
    _guardar(client)
    reg = BitacoraAuditoria.query.filter_by(accion='CONFIG_INSTITUCION').all()
    assert len(reg) == 1
    assert 'Instituto QA-7 del Valle' in reg[0].detalle and 'Mi Institución Educativa' in reg[0].detalle and 'Trimestre' in reg[0].detalle


@pytest.mark.parametrize('rol', [RolUsuario.ADMINISTRATIVO, RolUsuario.CONTADOR, RolUsuario.CAPTURADOR])
def test_solo_direccion_cambia_la_configuracion_de_la_institucion(client, app, rol):
    crear_usuario(username='otro', rol=rol); login(client, 'otro', 'clave12345')
    _guardar(client)
    assert _config().nombre_institucion != 'Instituto QA-7 del Valle'


def test_los_correos_a_alumnos_llevan_el_nombre_de_la_institucion(client, escenario):
    _guardar(client)
    client.application.config['MAIL_USERNAME'] = 'control@ejemplo.mx'
    with mail.record_messages() as bandeja:
        client.post(f"/cobro/{escenario['cargo'].id}/pagar", data={'monto_pagado': '100', 'metodo_pago': 'EFECTIVO'})
    assert len(bandeja) == 1
    assert 'Instituto QA-7 del Valle' in bandeja[0].body
    assert bandeja[0].sender and 'Instituto QA-7 del Valle' in str(bandeja[0].sender)


@pytest.mark.parametrize('url', ['/reportes/cartera-vencida/exportar', '/reportes/cobros-del-dia/exportar', '/cobros/dashboard/exportar'])
def test_las_exportaciones_de_excel_llevan_el_nombre_de_la_institucion(client, escenario, url):
    _guardar(client)
    ws = openpyxl.load_workbook(io.BytesIO(client.get(url).data)).worksheets[0]
    assert 'Instituto QA-7 del Valle' in str(ws['A1'].value)


def test_un_nombre_que_parece_formula_no_se_convierte_en_formula_en_las_exportaciones(client, escenario):
    _guardar(client, nombre_institucion='=HYPERLINK("http://x.mx","Uni")')
    ws = openpyxl.load_workbook(io.BytesIO(client.get('/reportes/cartera-vencida/exportar').data)).worksheets[0]
    assert ws['A1'].data_type != 'f'


# ------------------------------- 6. matrícula (ya cubierta a fondo en test_matricula_configurable) y persistencia -------------------------------

def test_cada_concepto_persiste_tras_recargar_y_no_afecta_a_los_demas(client, escenario):
    _guardar(client)
    c = _config()
    assert (c.nombre_institucion, c.nombre_periodo_singular, c.nombre_periodo_plural, c.nombre_programa_singular,
            c.nombre_programa_plural, c.max_periodos, c.matricula_digitos) == \
        ('Instituto QA-7 del Valle', 'Trimestre', 'Trimestres', 'Licenciatura', 'Licenciaturas', 7, 5)
    _guardar(client, nombre_institucion='Otro Nombre')                     # cambiar UNO no altera los demás
    c = _config()
    assert (c.nombre_institucion, c.nombre_periodo_singular, c.max_periodos) == ('Otro Nombre', 'Trimestre', 7)
