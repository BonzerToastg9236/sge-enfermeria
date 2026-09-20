"""
Auditoría 2026-09-19 (importación masiva):
  * los alumnos ACTIVOS importados quedaban sin historial, sin fecha de
    validación, sin cargos ni carga académica (y al reingresar desde baja se
    les cobraba "Inscripción" a alguien ya inscrito);
  * se aceptaban nombres que Excel interpreta como fórmula y valores más
    largos que la columna (500 en PostgreSQL a mitad del lote);
  * un xlsx de 5 KB con una celda lejana consumía ~420 MB;
  * "nan" en una calificación daba 500.
"""

import io
import zipfile
from decimal import Decimal
from unittest import mock

import pytest
from openpyxl import Workbook
from sqlalchemy.exc import DataError

from tests.conftest import crear_plan, crear_materia, crear_usuario, login
from app import db, Alumno, Cargo, ConceptoCobro, EstatusAlumno, InscripcionMateria, Calificacion
from modelos import HistorialEstatus

COLUMNAS = [
    'nombre_completo', 'curp', 'fecha_nacimiento', 'fecha_certificado_prepa', 'clave_carrera', 'sexo',
    'estatus', 'cuatrimestre_actual', 'telefono',
]


def _fila(i, nombre=None, estatus='ACTIVO', cuatri=1, telefono=None):
    return [nombre or f'Alumno Numero {i}', f'MASI{i:014d}', '2005-01-01', '2023-07-01', 'LEN', 'Femenino', estatus, cuatri, telefono]


def _xlsx(filas, extra=None):
    wb = Workbook(); ws = wb.active
    ws.append(COLUMNAS)
    for f in filas:
        ws.append(f)
        for celda in ws[ws.max_row]:
            if isinstance(celda.value, str) and celda.value.startswith('='):
                celda.data_type = 's'   # texto literal, como al teclear '=... en Excel
    if extra:
        extra(ws)
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return b


def _importar(client, archivo, generar_cargos=None):
    data = {'archivo_excel': (archivo, 'alumnos.xlsx')}
    if generar_cargos:
        data['generar_cargos'] = 'on'
    return client.post('/alumnos/importar', data=data, content_type='multipart/form-data', follow_redirects=True)


@pytest.fixture
def escenario(client, app):
    plan = crear_plan(clave='LEN'); plan.monto_mensualidad = Decimal('2500.00'); db.session.commit()
    db.session.add_all([
        ConceptoCobro(nombre='Inscripción', monto_sugerido=Decimal('3000.00'), activo=True),
        ConceptoCobro(nombre='Reinscripción', monto_sugerido=Decimal('2000.00'), activo=True),
        ConceptoCobro(nombre='Mensualidad', es_mensualidad=True, activo=True),
    ])
    db.session.commit()
    directivo = crear_usuario(); login(client, 'directivo1', 'clave12345')
    return plan, directivo


# --------------------------- historial / cargos ---------------------------

def test_los_alumnos_importados_como_activos_quedan_validados_y_con_historial(client, escenario):
    _, directivo = escenario
    _importar(client, _xlsx([_fila(1)]))
    alumno = Alumno.query.one()
    assert alumno.fecha_validacion is not None
    h = HistorialEstatus.query.one()
    assert (h.estatus_anterior, h.estatus_nuevo, h.usuario_fk) == (None, EstatusAlumno.ACTIVO, directivo.id)
    assert 'importaci' in h.comentario.lower()


def test_con_la_casilla_marcada_se_generan_inscripcion_mensualidades_y_carga(client, escenario):
    plan, _ = escenario
    crear_materia(plan, 'Anatomia', cuatrimestre=1)
    _importar(client, _xlsx([_fila(1)]), generar_cargos=True)
    conceptos = sorted(c.concepto for c in Cargo.query.all())
    assert conceptos.count('Inscripción') == 1 and conceptos.count('Mensualidad') == 4
    assert InscripcionMateria.query.count() == 1


def test_un_alumno_que_entra_a_cuatrimestre_avanzado_recibe_reinscripcion_no_inscripcion(client, escenario):
    _importar(client, _xlsx([_fila(1, cuatri=5)]), generar_cargos=True)
    conceptos = {c.concepto for c in Cargo.query.all()}
    assert 'Reinscripción' in conceptos and 'Inscripción' not in conceptos


def test_sin_la_casilla_no_se_generan_cargos_y_el_reingreso_no_cobra_inscripcion(client, escenario):
    _importar(client, _xlsx([_fila(1)]))
    assert Cargo.query.count() == 0
    a = Alumno.query.one()
    client.post(f'/alumno/{a.matricula_id}/cambiar-estatus', data={'nuevo_estatus': 'BAJA_TEMPORAL'})
    client.post(f'/alumno/{a.matricula_id}/cambiar-estatus', data={'nuevo_estatus': 'ACTIVO'})
    assert 'Inscripción' not in {c.concepto for c in Cargo.query.all()}


def test_un_alumno_importado_como_pendiente_no_se_valida_ni_se_le_cobra(client, escenario):
    _importar(client, _xlsx([_fila(1, estatus='PENDIENTE')]), generar_cargos=True)
    a = Alumno.query.one()
    assert a.fecha_validacion is None and Cargo.query.count() == 0


# --------------------------- datos hostiles ---------------------------

@pytest.mark.parametrize('nombre', ['=HYPERLINK("http://evil.example","x")', '+cmd|calc', '-2+3+cmd', '@SUM(A1)'])
def test_un_nombre_que_excel_leeria_como_formula_se_rechaza(client, escenario, nombre):
    respuesta = _importar(client, _xlsx([_fila(1, nombre=nombre)]))
    assert Alumno.query.count() == 0
    assert 'fórmula' in respuesta.get_data(as_text=True)


def test_un_valor_mas_largo_que_su_columna_se_reporta_en_la_fila(client, escenario):
    respuesta = _importar(client, _xlsx([_fila(1, telefono='5' * 25), _fila(2, nombre='X' * 201), _fila(3)]))
    assert Alumno.query.count() == 1
    html = respuesta.get_data(as_text=True)
    assert 'telefono' in html and 'nombre_completo' in html


def test_un_error_de_base_de_datos_en_una_fila_no_aborta_el_resto_del_lote(client, escenario):
    import rutas.alumnos as ruta
    real = ruta.crear_alumno_generando_matricula
    llamadas = []

    def falla_en_la_segunda(plan, **kw):
        llamadas.append(1)
        if len(llamadas) == 2:
            raise DataError('INSERT', {}, Exception('value too long for type character varying(20)'))
        return real(plan, **kw)

    with mock.patch.object(ruta, 'crear_alumno_generando_matricula', side_effect=falla_en_la_segunda):
        respuesta = _importar(client, _xlsx([_fila(1), _fila(2), _fila(3)]))
    assert respuesta.status_code == 200
    assert Alumno.query.count() == 2
    html = respuesta.get_data(as_text=True)
    assert '<td>3</td>' in html                       # la fila 3 del Excel (2ª de datos) figura entre los errores
    assert 'la base de datos rechazó esta fila' in html


# --------------------------- DoS por archivo ---------------------------

def test_una_celda_lejana_no_hace_que_se_recorran_cientos_de_miles_de_filas(client, escenario):
    def lejana(ws): ws.cell(row=100_000, column=5, value='x')
    respuesta = _importar(client, _xlsx([_fila(1)], extra=lejana))
    assert respuesta.status_code == 200
    assert Alumno.query.count() == 0
    assert 'demasiadas filas' in respuesta.get_data(as_text=True).lower()


def test_un_archivo_que_se_descomprime_a_cientos_de_megas_se_rechaza(client, escenario):
    base = _xlsx([_fila(1)]).getvalue()
    salida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as z_in, zipfile.ZipFile(salida, 'w', zipfile.ZIP_DEFLATED) as z_out:
        for item in z_in.infolist():
            z_out.writestr(item, z_in.read(item.filename))
        z_out.writestr('relleno.bin', b'\0' * (120 * 1024 * 1024))   # comprime a ~120 KB
    salida.seek(0)
    assert len(salida.getvalue()) < 1024 * 1024
    respuesta = _importar(client, salida)
    assert Alumno.query.count() == 0
    assert 'demasiado grande' in respuesta.get_data(as_text=True).lower()


def test_importar_boletas_tambien_rechaza_la_celda_lejana(client, app):
    plan = crear_plan(); crear_materia(plan, 'Anatomia', 1)
    from tests.conftest import crear_alumno
    crear_alumno(plan)
    crear_usuario(username='cap1', rol=__import__('app').RolUsuario.CAPTURADOR); login(client, 'cap1', 'clave12345')
    wb = Workbook(); ws = wb.active; ws.title = 'Calificaciones'
    ws.append(['Matrícula', 'Nombre Completo', 'Anatomia']); ws.cell(row=100_000, column=3, value='x')
    b = io.BytesIO(); wb.save(b); b.seek(0)
    r = client.post('/boletas/importar', data={'plan_id': plan.id, 'cuatrimestre': 1, 'periodo_escolar': '2026-B',
                                               'archivo_excel': (b, 'b.xlsx')}, content_type='multipart/form-data')
    assert r.status_code == 200 and 'demasiadas filas' in r.get_data(as_text=True).lower()


# --------------------------- calificaciones ---------------------------

@pytest.mark.parametrize('valor', ['nan', 'NaN', 'inf', '-inf', '1e1', '10.0001', '-0.5', '8,5'])
def test_una_calificacion_no_finita_o_rara_no_se_guarda_ni_da_500(client, app, valor):
    plan = crear_plan(); materia = crear_materia(plan, 'Anatomia', 1)
    from tests.conftest import crear_alumno
    a = crear_alumno(plan)
    crear_usuario(username='cap1', rol=__import__('app').RolUsuario.CAPTURADOR); login(client, 'cap1', 'clave12345')
    r = client.post(f'/alumno/{a.matricula_id}/boleta', data={'periodo_escolar': '2026-B', 'cuatrimestre': '1', f'calificacion_{materia.id}': valor})
    assert r.status_code == 302
    assert Calificacion.query.count() == 0


def test_una_calificacion_valida_se_sigue_guardando(client, app):
    plan = crear_plan(); materia = crear_materia(plan, 'Anatomia', 1)
    from tests.conftest import crear_alumno
    a = crear_alumno(plan)
    crear_usuario(username='cap1', rol=__import__('app').RolUsuario.CAPTURADOR); login(client, 'cap1', 'clave12345')
    client.post(f'/alumno/{a.matricula_id}/boleta', data={'periodo_escolar': '2026-B', 'cuatrimestre': '1', f'calificacion_{materia.id}': '9.5'})
    assert Calificacion.query.one().calificacion_final == 9.5
