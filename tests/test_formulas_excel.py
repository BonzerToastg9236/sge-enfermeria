"""
Pruebas anti-regresión de la inyección de fórmulas en Excel (hallazgo #2
de la auditoría, plan de remediación paso 2).

CONTEXTO: rutas/academico.py::plantilla_boletas() escribe
alumno.nombre_completo directo en una celda .xlsx. Si el texto empieza
con = + - @, Excel/Sheets lo interpreta como fórmula al abrir el
archivo -- y el nombre viene de /registro, el único formulario público
sin login.

Dos capas de defensa, cada una con su propia prueba:
1. rutas/registro.py valida el nombre con NOMBRE_REGEX (evita que un
   nombre así llegue a existir por esa vía).
2. utilidades/archivos.py::valor_seguro_excel() antepone un apóstrofo a
   cualquier texto que empiece con esos caracteres, sin importar de
   dónde venga (datos ya existentes, importación masiva, etc.).
"""

import io

from openpyxl import load_workbook

from tests.conftest import crear_plan, crear_materia, crear_alumno, crear_usuario, login

from app import RolUsuario
from utilidades.archivos import valor_seguro_excel


# ---------------------------------------------------------------------------
# Capa 2: utilidades/archivos.py::valor_seguro_excel()
# ---------------------------------------------------------------------------

def test_valor_seguro_excel_antepone_apostrofo_si_empieza_con_igual():
    assert valor_seguro_excel('=cmd|"/c calc"!A1') == '\'=cmd|"/c calc"!A1'


def test_valor_seguro_excel_antepone_apostrofo_para_los_cuatro_caracteres_de_formula():
    for caracter in ('=', '+', '-', '@'):
        valor = f'{caracter}1+1'
        assert valor_seguro_excel(valor) == f"'{valor}", (
            f'valor_seguro_excel no neutralizó un valor que empieza con "{caracter}"'
        )


def test_valor_seguro_excel_no_modifica_texto_normal():
    assert valor_seguro_excel('María Fernanda López') == 'María Fernanda López'


def test_valor_seguro_excel_no_falla_con_valores_no_string():
    # Las celdas de matrícula/otras columnas pueden traer None o números;
    # el helper no debe intentar indexar algo que no sea texto.
    assert valor_seguro_excel(None) is None
    assert valor_seguro_excel(42) == 42


# ---------------------------------------------------------------------------
# Capa 1: rutas/registro.py -- NOMBRE_REGEX
# ---------------------------------------------------------------------------

def test_registro_rechaza_nombre_que_empieza_con_signo_de_formula(client, app):
    from tests.test_registro_publico import _datos_con_plan

    plan = crear_plan()

    respuesta = client.post(
        '/registro',
        data=_datos_con_plan(plan.id, nombre_completo='=cmd|"/c calc"!A1'),
        follow_redirects=True
    )

    assert respuesta.status_code == 400
    assert 'nombre'.encode() in respuesta.data.lower(), (
        'Debería rechazar el nombre con un mensaje de error, no aceptarlo'
    )

    from app import Alumno
    assert Alumno.query.filter_by(curp='LORF050101MDFXYZ09').first() is None, (
        'No debería haberse creado el alumno con un nombre inválido'
    )


def test_registro_acepta_nombre_normal_con_acentos(client, app):
    from tests.test_registro_publico import _datos_con_plan

    plan = crear_plan()

    respuesta = client.post(
        '/registro',
        data=_datos_con_plan(plan.id, nombre_completo='María Fernanda López Ramírez'),
        follow_redirects=True
    )

    assert respuesta.status_code == 200
    from app import Alumno
    assert Alumno.query.filter_by(curp='LORF050101MDFXYZ09').first() is not None


def test_registro_acepta_nombre_con_guion_y_apostrofo(client, app):
    """Apellidos compuestos reales: 'Pérez-García', 'D'León'."""
    from tests.test_registro_publico import _datos_con_plan

    plan = crear_plan()

    respuesta = client.post(
        '/registro',
        data=_datos_con_plan(plan.id, nombre_completo="Ana Pérez-García D'León"),
        follow_redirects=True
    )

    assert respuesta.status_code == 200
    from app import Alumno
    assert Alumno.query.filter_by(curp='LORF050101MDFXYZ09').first() is not None


def test_registro_sigue_rechazando_nombre_que_empieza_con_guion(client, app):
    """
    Permitir '-' dentro del nombre (apellidos compuestos) no debe reabrir
    la inyección de fórmulas: un nombre que EMPIEZA con '-' sigue
    prohibido, igual que con '=', '+' y '@'.
    """
    from tests.test_registro_publico import _datos_con_plan

    plan = crear_plan()

    respuesta = client.post(
        '/registro',
        data=_datos_con_plan(plan.id, nombre_completo='-1+1'),
        follow_redirects=True
    )

    assert respuesta.status_code == 400
    from app import Alumno
    assert Alumno.query.filter_by(curp='LORF050101MDFXYZ09').first() is None


# ---------------------------------------------------------------------------
# Extremo a extremo: la plantilla exportada neutraliza cualquier nombre
# hostil que exista en la base (importación masiva, dato legado, etc.),
# independientemente de si pasó por /registro.
# ---------------------------------------------------------------------------

def test_plantilla_boletas_neutraliza_nombre_con_formula(client, app):
    plan = crear_plan()
    crear_materia(plan, nombre='Anatomía', cuatrimestre=1)
    crear_alumno(plan, nombre='=cmd|"/c calc"!A1', curp='FORM010101HDFXYZ01')
    crear_usuario(rol=RolUsuario.DIRECTIVO)
    login(client, 'directivo1', 'clave12345')

    respuesta = client.get(f'/boletas/importar/plantilla?plan_id={plan.id}&cuatrimestre=1')

    assert respuesta.status_code == 200
    libro = load_workbook(io.BytesIO(respuesta.data))
    hoja = libro['Calificaciones']
    valor_nombre = hoja.cell(row=2, column=2).value

    assert valor_nombre.startswith("'"), (
        f'La celda del nombre debería quedar neutralizada con un apóstrofo inicial, salió: {valor_nombre!r}'
    )
